"""Scene-specific class-meeting lesson / parent-meeting implementation plan.

The model writes the substance. This module validates alignment and renders it;
it does not synthesize a lesson from slide titles or pad a textbook template.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from .deliverable_naming import project_deliverable_relpaths
from .document_rendering import (find_pdf_font_path, p, pdf_page_count, pdf_text_probe,
                                 require_locked_source_exists, require_locked_source_matches_state)
from .events import append_event
from .file_hashing import file_sha256
from .json_io import read_json, write_json
from .state import read_state, write_state
from .thematic_workflow import education_context_from_brief, require_content_review, require_step
from .time_utils import now_iso
from .validation import ValidationError, validate_lesson_plan_manifest


def _texts(value: Any, label: str) -> list[str]:
    if not isinstance(value, list) or not value or any(not isinstance(s, str) or not s.strip() for s in value):
        raise ValidationError(f"{label} requires concrete, non-empty text items")
    return value


def validate_thematic_plan(root: Path, plan: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(plan, dict) or plan.get("schema_version") != "thematic-1.0":
        raise ValidationError("thematic plan requires schema_version=thematic-1.0; use the scene-specific template")
    if plan.get("project_name") != state["project_name"] or plan.get("run_dir") != state["run_dir"]:
        raise ValidationError("thematic plan must belong to this project")
    brief = state["thematic_brief"]
    for key in ("scene", "school_stage", "grade", "duration_minutes"):
        if plan.get(key) != brief.get(key):
            raise ValidationError(f"thematic plan {key} differs from the current target")
    if not isinstance(plan.get("theme"), str) or not plan["theme"].strip():
        raise ValidationError("theme is required")
    basis = plan.get("basis", {})
    source = basis.get("locked_presentation_source")
    if not isinstance(source, dict):
        raise ValidationError("actual locked presentation source is required")
    require_locked_source_exists(root, source)
    require_locked_source_matches_state(state, source)
    for key in ("theme_analysis", "audience_analysis", "goals", "key_points", "difficult_points", "preparation", "follow_up", "reflection_prompts"):
        _texts(plan.get(key), key)
    activities = plan.get("activities")
    if not isinstance(activities, list) or not activities:
        raise ValidationError("activities are required")
    page_titles = {s["slide_index"]: s["title"] for s in read_json(root / "_state/阶段1/content.json")["slides"]}
    covered = set()
    duration = 0
    for activity in activities:
        for key in ("name", "purpose"):
            if not isinstance(activity.get(key), str) or not activity[key].strip():
                raise ValidationError(f"activity.{key} is required")
        minutes = activity.get("duration_minutes")
        if type(minutes) is not int or minutes <= 0:
            raise ValidationError("each activity duration must be a positive integer")
        duration += minutes
        refs = activity.get("pages")
        if not isinstance(refs, list) or not refs:
            raise ValidationError("activities need actual PPT page/title references")
        for ref in refs:
            if not isinstance(ref, dict) or type(ref.get("slide_index")) is not int or page_titles.get(ref["slide_index"]) != ref.get("title"):
                raise ValidationError("activity page/title reference is invalid")
            covered.add(ref["slide_index"])
        for key in ("teacher_actions", "participant_actions", "materials", "output", "assessment", "possible_responses", "followup_questions", "fallback"):
            _texts(activity.get(key), f"activity.{key}")
    if duration != brief["duration_minutes"]:
        raise ValidationError("main activity durations must total the specified session duration; alternatives replace existing activities")
    if covered != set(page_titles):
        raise ValidationError("the implementation plan must account for every deck page")
    if not isinstance(plan.get("applicability"), str) or not plan["applicability"].strip():
        raise ValidationError("honest main applicable range is required")
    adaptations = plan.get("grade_adaptations")
    if not isinstance(adaptations, list):
        raise ValidationError("grade_adaptations must be a list of short replacement approaches")
    if brief.get("grade") == "通用版" and not adaptations:
        raise ValidationError("general-band version needs short simplification/extension alternatives, not separate decks")
    for item in adaptations:
        if item.get("activity") not in {a["name"] for a in activities}:
            raise ValidationError("adaptations must refer to a real activity")
        for key in ("scope", "replacement", "time_rule"):
            if not isinstance(item.get(key), str) or not item[key].strip():
                raise ValidationError(f"grade adaptation {key} is required")
    return plan


def plan_sections(plan: dict[str, Any]) -> list[tuple[str, list[str]]]:
    parent = plan["scene"] == "parent_meeting"
    participant = "家长" if parent else "学生"
    result = [
        ("使用说明", [f"适用学段：{ {'primary':'小学','junior_high':'初中','senior_high':'高中'}[plan['school_stage']]}｜{plan['grade']}",
                      f"建议时长：{plan['duration_minutes']}分钟", f"主要适用范围：{plan['applicability']}"]),
        ("主题与真实使用情境", plan["theme_analysis"]),
        ("家长需求与沟通难点" if parent else "学情与理解障碍", plan["audience_analysis"]),
        ("会议目标与家校共识" if parent else "班会学习与行动目标", plan["goals"]),
        ("核心问题", plan["key_points"]), ("难点与突破", plan["difficult_points"]),
        ("会前准备" if parent else "课前准备", plan["preparation"]),
    ]
    for a in plan["activities"]:
        refs = "、".join(f"第{r['slide_index']}页《{r['title']}》" for r in a["pages"])
        lines = [f"时间：{a['duration_minutes']}分钟｜页面：{refs}", f"目标：{a['purpose']}"]
        for label, key in (("教师组织", "teacher_actions"), (f"{participant}任务", "participant_actions"),
                           ("材料与使用方式", "materials"), ("预期任务产出（非真实效果）", "output"),
                           ("观察与评价", "assessment"), ("可能回应（供备课参考）", "possible_responses"),
                           ("追问与引导", "followup_questions"), ("冷场／时间不足备选", "fallback")):
            lines += [f"{label}：{s}" for s in a[key]]
        result.append((a["name"], lines))
    alternatives = [f"{a['scope']}｜{a['activity']}：{a['replacement']}（{a['time_rule']}）" for a in plan["grade_adaptations"]]
    if alternatives:
        result.append(("年级适配提示：替换原活动，不叠加时长", alternatives))
    result += [("会后家庭行动与反馈" if parent else "课后行动与迁移", plan["follow_up"]),
               ("实施后复盘提示（授课前不填写真实效果）", plan["reflection_prompts"])]
    return result


def build_thematic_plan(run_dir: str | Path, plan_json: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    if state["current_stage"] != "stage3" or state["required_actor"] != "main_controller":
        raise ValidationError("theme plan is only rendered in stage3 after full-deck acceptance")
    require_step(state, 3)
    require_content_review(root, "deck")
    plan = validate_thematic_plan(root, read_json(plan_json), state)
    name = "会议实施方案" if plan["scene"] == "parent_meeting" else "班会教案"
    title = f"{plan['theme']}｜{name}"
    sections = plan_sections(plan)
    relpaths = project_deliverable_relpaths(root, state=state)
    files = []
    for ext, key in (("md", "stage3_lesson_plan"), ("docx", "stage3_lesson_plan_docx"), ("pdf", "stage3_lesson_plan_pdf")):
        path = root / relpaths[key]
        path.parent.mkdir(parents=True, exist_ok=True)
        if ext == "md":
            path.write_text("# " + title + "\n\n" + "\n\n".join("## " + h + "\n\n" + "\n".join("- " + s for s in items) for h, items in sections) + "\n", encoding="utf-8")
        elif ext == "docx":
            _render_docx(title, sections, path)
        else:
            _render_pdf(title, sections, path)
        files.append({"label": f"{ext.upper()} {name}", "path": relpaths[key], "sha256": file_sha256(path)})
        state["user_artifacts"][key] = relpaths[key]
        state["expected_user_paths"][key] = relpaths[key]
    write_json(root / "_state/阶段3/lesson_plan.json", {**plan, "education_context": education_context_from_brief(state["thematic_brief"])})
    manifest = validate_lesson_plan_manifest({"schema_version": "1.0", "project_name": state["project_name"],
        "run_dir": state["run_dir"], "files": files, "status": "generated", "created_at": now_iso(),
        "conversion": {"source": "_state/阶段3/lesson_plan.json", "pdf_path": relpaths["stage3_lesson_plan_pdf"],
                       "paired_docx": relpaths["stage3_lesson_plan_docx"], "tool": "reportlab_scene_plan"},
        "summary": {"periods": 1, "activities": len(plan["activities"]), "subject_group": "integrated",
                    "word_table_mode": "scene_specific_activity_sections", "plan_type": name,
                    "pdf_pages": pdf_page_count(root / relpaths["stage3_lesson_plan_pdf"]),
                    "pdf_text_probe": pdf_text_probe(root / relpaths["stage3_lesson_plan_pdf"])}})
    write_json(root / "_state/阶段3/lesson_plan_manifest.json", manifest)
    state["stage3_outputs"]["lesson_plan"].update({"required": True, "status": "generated"})
    state["status"] = "stage3_lesson_plan_generated"
    state["next_required_action"] = "完成对应逐字稿并实际审阅配套内容与排版，然后连续整理交付"
    write_state(root, state)
    append_event(root, "thematic_plan_generated", "runtime", plan_type=name)
    return manifest


def _render_docx(title: str, sections: list, path: Path) -> None:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.35
    doc.add_heading(title, 0)
    for heading, items in sections:
        doc.add_heading(heading, 1)
        for text in items:
            doc.add_paragraph(text, style="List Bullet")
    from .document_rendering import add_page_number
    add_page_number(sec.footer.paragraphs[0])
    doc.save(path)


def _render_pdf(title: str, sections: list, path: Path) -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    pdfmetrics.registerFont(TTFont("ThematicCJK", str(find_pdf_font_path()), subfontIndex=0))
    body = ParagraphStyle("ThemeBody", fontName="ThematicCJK", fontSize=10.5, leading=16, spaceAfter=6, wordWrap="CJK")
    heading = ParagraphStyle("ThemeHeading", parent=body, fontSize=13, leading=20, spaceBefore=12, keepWithNext=True)
    title_style = ParagraphStyle("ThemeTitle", parent=heading, fontSize=18, leading=26)
    story = [Paragraph(p(title), title_style), Spacer(1, 8)]
    for h, items in sections:
        story.append(Paragraph(p(h), heading))
        story += [Paragraph(p("• " + s), body) for s in items]
    def footer(canvas, doc):
        canvas.setFont("ThematicCJK", 8)
        canvas.drawCentredString(A4[0] / 2, 24, f"第{doc.page}页")
    SimpleDocTemplate(str(path), pagesize=A4, topMargin=45, bottomMargin=42,
                      leftMargin=48, rightMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
