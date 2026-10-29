"""File-only fade timing and hidden-window rendering. Never start a slideshow.

The supported timing grammar is deliberately narrow. Unknown playback behavior
raises an error instead of producing a plausible but unverified state image.
"""
from __future__ import annotations

from copy import deepcopy
import io
import json
import math
import os
from pathlib import Path
import posixpath
import shutil
import struct
import subprocess
import tempfile
import uuid
import zipfile

from lxml import etree

from .file_hashing import file_sha256
from .json_io import read_json, write_json
from .time_utils import now_iso
from .validation import ValidationError

P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
REL = '{http://schemas.openxmlformats.org/package/2006/relationships}'
CT = '{http://schemas.openxmlformats.org/package/2006/content-types}'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
PROFILE = 'ooxml_fade_v1'
RENDERER = 'PowerPoint hidden Slide.Export + temporary embedded fonts'


def _parts(path):
    from .teaching_animation import slide_parts
    return slide_parts(path)


def _top_shapes(tree):
    shapes = tree.find(f'{P}cSld/{P}spTree')
    result = {}
    for node in shapes:
        if etree.QName(node).localname not in {'sp','grpSp','graphicFrame','cxnSp','pic','contentPart'}:
            continue
        nv = next(node.iter(P+'cNvPr'), None)
        if nv is not None:
            result[int(nv.get('id'))] = node
    return result


def presentation_uses_timings(package):
    """Read the related presentation properties; unknown settings stay unknown."""
    try:
        rels = etree.fromstring(package.read('ppt/_rels/presentation.xml.rels'))
        props = [r for r in rels if r.get('Type') == R[1:-1] + '/presProps']
        if len(props) != 1 or props[0].get('TargetMode') == 'External':
            return None
        target = props[0].get('Target', '')
        part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('ppt', target))
        tree = etree.fromstring(package.read(part))
        shows = tree.findall(P + 'showPr')
        if len(shows) != 1:
            return None
        value = shows[0].get('useTimings')
        return {'true': True, '1': True, 'false': False, '0': False}.get(value)
    except (KeyError, etree.XMLSyntaxError):
        return None


def inspect_slide_advance(tree, use_timings):
    # Check both markup-compatibility branches conservatively. Do not guess
    # which Office version will select a branch or rewrite its transition.
    mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
    transitions = tree.findall(P + 'transition') + tree.findall(mc + 'AlternateContent//' + P + 'transition')
    return {'advance_after': [t.get('advTm') for t in transitions if t.get('advTm') is not None],
            'use_timings': use_timings}


def require_manual_advance(advance, page):
    if not isinstance(advance, dict) or not isinstance(advance.get('advance_after'), list):
        raise ValidationError(f'第{page}页缺少自动翻页检查依据；请重新检查当前输入')
    if advance['advance_after'] and advance.get('use_timings') is not False:
        setting = '启用计时' if advance.get('use_timings') is True else '未能确认计时已关闭'
        raise ValidationError(f'第{page}页存在自动翻页（advTm={advance["advance_after"]}，{setting}），'
                              '与人工点击揭示冲突；须单独处理放映设置，原切换保持不变')


def parse_fade_timing(tree):
    """Read final OOXML, not the proposed plan. Return settled click states."""
    timing = tree.find(P+'timing')
    shapes = _top_shapes(tree)
    if timing is None:
        return []
    clock_ids = [n.get('id', '') for n in timing.iter(P + 'cTn')]
    if any(not value or not value.isascii() or not value.isdecimal() or int(value) > 4294967295
           for value in clock_ids) or len({int(value) for value in clock_ids}) != len(clock_ids):
        raise ValidationError('动画时间节点编号须为有效且不重复的无符号32位整数')
    if tree.findall('.//'+P+'audio') or tree.findall('.//'+P+'video'):
        raise ValidationError('后台时间轴不支持音视频播放；不得以静态状态验收')
    if timing.find('.//'+P+'seq') is None:
        raise ValidationError('动画时间轴缺少主序列')
    seqs = timing.findall('.//'+P+'seq')
    if len(seqs) != 1 or len(timing.findall('.//'+P+'cTn')) == 0:
        raise ValidationError('后台时间轴不支持交互或多主序列')
    seq = seqs[0]
    if seq.get('concurrent') != '1' or seq.get('nextAc') != 'seek':
        raise ValidationError('未支持的动画序列语义')
    main = seq.find(P+'cTn')
    if main is None or main.get('nodeType') != 'mainSeq':
        raise ValidationError('未支持的动画主节点')
    # Effects other than fade entrance, motion, paragraph builds, repeat,
    # restart, target triggers and dynamic subnodes must not be approximated.
    allowed_tags = {'timing','tnLst','par','seq','cTn','childTnLst','stCondLst',
        'cond','set','cBhvr','tgtEl','spTgt','attrNameLst','attrName','to','strVal',
        'animEffect','prevCondLst','nextCondLst','sldTgt','bldLst','bldP'}
    for node in timing.iter():
        if etree.QName(node).namespace != P[1:-1] or etree.QName(node).localname not in allowed_tags:
            raise ValidationError('未支持的动画节点: '+etree.QName(node).localname)
        for key in ('repeatCount','repeatDur','autoRev','spd','accel','decel'):
            if key in node.attrib:
                raise ValidationError('未支持的动画重复或速度变换')
        if node.tag == P+'spTgt' and (list(node) or set(node.attrib) != {'spid'}):
            raise ValidationError('未支持的分段文字动画')
        if node.tag == P+'bldP' and set(node.attrib)-{'spid','grpId'}:
            raise ValidationError('未支持的文字构建方式')
    for parent, event in [('prevCondLst','onPrev'), ('nextCondLst','onNext')]:
        conds=seq.findall(f'{P}{parent}/{P}cond')
        if len(conds)!=1 or conds[0].attrib!={'evt':event,'delay':'0'} or conds[0].find(f'{P}tgtEl/{P}sldTgt') is None:
            raise ValidationError('未支持的点击推进条件')
    effect_nodes=timing.xpath('.//p:cTn[@presetClass]',namespaces={'p':P[1:-1]})
    behavior_targets=[]; effects=[]; seen=set(); click=0
    for node in effect_nodes:
        trigger={'clickEffect':'click','withEffect':'with_previous','afterEffect':'after_previous'}.get(node.get('nodeType'))
        if trigger is None or node.get('presetClass')!='entr' or node.get('presetID')!='10' or node.get('fill')!='hold' or node.get('presetSubtype','0')!='0':
            raise ValidationError('后台仅支持完整对象短淡入；未知效果不能验收')
        conditions=node.findall(f'{P}stCondLst/{P}cond')
        if len(conditions)!=1 or conditions[0].attrib!={'delay':'0'}:
            raise ValidationError('未支持的动画延时或触发条件')
        if trigger=='click': click+=1
        if click==0:
            raise ValidationError('自动开始动画没有明确的初始点击边界')
        children=node.find(P+'childTnLst')
        if children is None or [etree.QName(n).localname for n in children]!=['set','animEffect']:
            raise ValidationError('淡入行为组合不完整或含未支持行为')
        aset, fade=children
        targets=node.findall('.//'+P+'spTgt')
        ids={int(n.get('spid')) for n in targets if n.get('spid','').isdigit()}
        if len(targets)!=2 or len(ids)!=1:
            raise ValidationError('淡入目标引用不一致')
        sid=next(iter(ids))
        if sid not in shapes or sid in seen:
            raise ValidationError('重复、嵌套或无效的动画目标')
        seen.add(sid);behavior_targets.extend(targets)
        if fade.attrib!={'transition':'in','filter':'fade'}:
            raise ValidationError('未支持的淡入滤镜')
        if [x.text for x in aset.findall('.//'+P+'attrName')]!=['style.visibility'] or aset.find(f'{P}to/{P}strVal').get('val')!='visible':
            raise ValidationError('淡入可见性行为不一致')
        duration=fade.find(f'{P}cBhvr/{P}cTn')
        if duration is None or not duration.get('dur','').isdigit():
            raise ValidationError('淡入时长无效')
        ms=int(duration.get('dur'))
        if not 200<=ms<=1200:
            raise ValidationError('后台支持的淡入时长为200—1200毫秒')
        effects.append({'shape_id':sid,'trigger':trigger,'click':click,'duration':ms/1000})
    if not effects or len(timing.findall('.//'+P+'spTgt'))!=len(behavior_targets):
        raise ValidationError('空时间轴或隐藏的未支持目标')
    build_ids=set()
    for build in timing.findall(f'{P}bldLst/{P}bldP'):
        sid=build.get('spid','')
        if not sid.isdigit() or int(sid) not in seen or int(sid) in build_ids or build.get('grpId')!='0':
            raise ValidationError('未支持或无效的文字构建目标/分组')
        build_ids.add(int(sid))
    # Validate structural grouping, not just document order: every top-level
    # mainSeq child is one click, every inner batch begins at delay zero.
    containers=main.findall(f'{P}childTnLst/{P}par/{P}cTn')
    if len(containers)!=click:
        raise ValidationError('XML点击容器与效果触发条件不一致')
    for number, container in enumerate(containers,1):
        cs=container.findall(f'{P}stCondLst/{P}cond')
        if len(cs)!=1 or cs[0].attrib!={'delay':'indefinite'}:
            raise ValidationError('点击容器没有等待点击')
        inner=container.findall(f'{P}childTnLst/{P}par/{P}cTn')
        if len(inner)!=1 or inner[0].find(f'{P}stCondLst/{P}cond').attrib!={'delay':'0'}:
            raise ValidationError('未支持的组内延时')
        contained=container.xpath('.//p:cTn[@presetClass]',namespaces={'p':P[1:-1]})
        expected=[e for e in effects if e['click']==number]
        if len(contained)!=len(expected) or any(n is not effect_nodes[sum(len([e for e in effects if e['click']==k]) for k in range(1,number))+i] for i,n in enumerate(contained)):
            raise ValidationError('动画XML分组顺序与点击不一致')
        if any(n.get('nodeType')=='afterEffect' for n in contained):
            # after_previous requires precise chained time containers. Support
            # it only when written/parsed by a future dedicated profile.
            raise ValidationError('当前后台支持click与with_previous；after_previous需明确改为同步组或列为阻断')
    # Reject any timing semantics not accounted for above. IDs and grpId are
    # bookkeeping; default whole-shape bldP nodes do not change visibility.
    def signature(node):
        attrs={k:v for k,v in node.attrib.items() if k not in ('id','grpId')}
        if node.tag==P+'cTn' and 'presetClass' in attrs:
            attrs.setdefault('presetSubtype','0')
        return (node.tag,tuple(sorted(attrs.items())),(node.text or '').strip(),
            tuple(signature(c) for c in node if c.tag!=P+'bldLst'))
    if signature(timing)!=signature(make_fade_timing(effects)):
        raise ValidationError('动画时间树含未建模条件或行为；不能以近似状态验收')
    return effects


def _el(parent, name, **attrs):
    return etree.SubElement(parent,P+name,{k:str(v) for k,v in attrs.items()})


def make_fade_timing(effects):
    if not effects: return None
    timing=etree.Element(P+'timing');counter=iter(range(1,1000000))
    def tn(parent,**attrs): return _el(parent,'cTn',id=next(counter),**attrs)
    root=tn(_el(_el(timing,'tnLst'),'par'),dur='indefinite',restart='never',nodeType='tmRoot')
    seq=_el(_el(root,'childTnLst'),'seq',concurrent='1',nextAc='seek')
    main=tn(seq,dur='indefinite',nodeType='mainSeq');ml=_el(main,'childTnLst')
    batch=None
    for effect in effects:
        trigger=effect['trigger']
        if trigger not in {'click','with_previous'} or (batch is None and trigger!='click'):
            raise ValidationError('后台淡入仅支持单击及同组同步；首效果必须单击')
        seconds=effect.get('duration',.45)
        if type(seconds) not in (int,float) or not math.isfinite(seconds) or not .2<=seconds<=1.2:
            raise ValidationError('无效淡入时长')
        if trigger=='click':
            outer=tn(_el(ml,'par'),fill='hold');_el(_el(outer,'stCondLst'),'cond',delay='indefinite')
            inner=tn(_el(_el(outer,'childTnLst'),'par'),fill='hold');_el(_el(inner,'stCondLst'),'cond',delay='0');batch=_el(inner,'childTnLst')
        et=tn(_el(batch,'par'),presetID='10',presetClass='entr',presetSubtype='0',fill='hold',nodeType='clickEffect' if trigger=='click' else 'withEffect')
        _el(_el(et,'stCondLst'),'cond',delay='0');cl=_el(et,'childTnLst')
        aset=_el(cl,'set');bh=_el(aset,'cBhvr');bt=tn(bh,dur='1',fill='hold');_el(_el(bt,'stCondLst'),'cond',delay='0')
        _el(_el(bh,'tgtEl'),'spTgt',spid=effect['shape_id']);_el(_el(bh,'attrNameLst'),'attrName').text='style.visibility';_el(_el(aset,'to'),'strVal',val='visible')
        fade=_el(cl,'animEffect',transition='in',filter='fade');bh=_el(fade,'cBhvr');tn(bh,dur=round(seconds*1000));_el(_el(bh,'tgtEl'),'spTgt',spid=effect['shape_id'])
    for name,event in [('prevCondLst','onPrev'),('nextCondLst','onNext')]:
        _el(_el(_el(seq,name),'cond',evt=event,delay='0'),'tgtEl');seq[-1][0][0].append(etree.Element(P+'sldTgt'))
    return timing


def write_timing_file(source,output,plan):
    parts=_parts(source); mapped={parts[s['page']-1]:s for s in plan['slides']}
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(output,'w') as b:
        use_timings = presentation_uses_timings(a)
        for item in a.infolist():
            data=a.read(item.filename);entry=mapped.get(item.filename)
            if entry and (entry['effects'] or entry.get('replace_targets')):
                tree=etree.fromstring(data)
                require_manual_advance(inspect_slide_advance(tree, use_timings), entry['page'])
                prior=parse_fade_timing(tree)
                top=list(_top_shapes(tree));replace={top[i-1] for i in entry.get('replace_targets',[])}
                retained=[e for e in prior if e['shape_id'] not in replace]
                # Retain click boundaries when the original click leader was
                # replaced, but other members of that click remain.
                last=None
                for e in retained:
                    e['trigger']='click' if e['click']!=last else 'with_previous';last=e['click']
                timing=make_fade_timing(retained+entry['effects']);old=tree.find(P+'timing')
                if old is not None:tree.remove(old)
                ext=tree.find(P+'extLst');tree.insert(tree.index(ext) if ext is not None else len(tree),timing)
                parse_fade_timing(tree)
                data=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
            b.writestr(item,data)


def read_click_states(path):
    pages=[]
    with zipfile.ZipFile(path) as z:
        use_timings = presentation_uses_timings(z)
        for page,part in enumerate(_parts(path),1):
            tree=etree.fromstring(z.read(part));effects=parse_fade_timing(tree);top=_top_shapes(tree)
            if effects:
                require_manual_advance(inspect_slide_advance(tree, use_timings), page)
            for tag in ('audio','video','oleObj','control','contentPart'):
                if tree.find('.//'+P+tag) is not None:
                    raise ValidationError(f'第{page}页包含未支持的动态对象{tag}')
            if tree.find('.//'+A+'fld') is not None:
                raise ValidationError(f'第{page}页包含动态字段，无法可靠复制每次点击画面')
            if any(tree.find('.//'+A+tag) is not None for tag in ('hlinkClick','hlinkMouseOver')):
                raise ValidationError(f'第{page}页包含点击/鼠标交互，不能当普通推进状态验收')
            if any(n.get('hidden') not in (None,'0','false') for n in tree.iter(P+'cNvPr')):
                raise ValidationError(f'第{page}页包含未支持的隐藏对象属性')
            reveal={e['shape_id']:e['click'] for e in effects};count=max(reveal.values(),default=0)
            states=[{'step':step,'visible_ids':[sid for sid in top if reveal.get(sid,0)<=step],
                     'hidden_ids':[sid for sid in top if reveal.get(sid,0)>step]} for step in range(count+1)]
            pages.append({'page':page,'part':part,'effects':effects,'states':states})
    return pages


def extract_embedded_fonts(source,directory):
    """Validated EOT/MTX; original glyphs, no font installation or substitution."""
    from fontTools.ttLib import TTFont
    paths=[];aliases={};suffix=uuid.uuid4().hex[:10]
    with zipfile.ZipFile(source) as z:
        pres=etree.fromstring(z.read('ppt/presentation.xml'))
        rels={x.get('Id'):x.get('Target') for x in etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))}
        records=[]
        for embedded in pres.findall('.//'+P+'embeddedFont'):
            face=embedded.find(P+'font').get('typeface')
            aliases.setdefault(face,'PptBg'+suffix+'F'+str(len(aliases)))
            for variant in embedded:
                rid=variant.get(R+'id')
                if rid:
                    target=rels[rid];name=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('ppt',target))
                    records.append((name,face,etree.QName(variant).localname))
        for name,face,variant in records:
            raw=z.read(name)
            from .embedded_fonts import decode_font
            data,_=decode_font(raw)
            stream=io.BytesIO(data)
            stream.name='<embedded-font-memory>'
            font=TTFont(stream, lazy=True, recalcBBoxes=False, recalcTimestamp=False)
            flags=font['OS/2'].fsType if 'OS/2' in font else 0
            if flags & (2|256|512):raise ValidationError('嵌入字体不允许当前临时渲染方式')
            # Some PDF converters put a full face name ("... Bold") in
            # typeface while the embedded font's family name is different.
            # Hidden export otherwise silently substitutes Songti. Alias only
            # the disposable font name table; all glyphs and metrics stay exact.
            alias=aliases[face]
            style={'regular':'Regular','bold':'Bold','italic':'Italic','boldItalic':'Bold Italic'}.get(variant,'Regular')
            for record in font['name'].names:
                if record.nameID in (1,4,6,16):
                    value=alias if record.nameID!=6 else alias+'-'+style.replace(' ','')
                    record.string=value.encode(record.getEncoding())
                elif record.nameID in (2,17):record.string=style.encode(record.getEncoding())
            target=Path(directory)/('font_'+str(len(paths))+'.ttf');font.save(target);paths.append(str(target))
            font.close()
    return paths,aliases


def prepare_font_render_copy(source,target,aliases):
    """Render-only names map to the very same embedded glyph data."""
    with zipfile.ZipFile(source) as z,zipfile.ZipFile(target,'w') as out:
        for item in z.infolist():
            data=z.read(item.filename)
            if item.filename.endswith('.xml'):
                tree=etree.fromstring(data);changed=False
                for node in tree.iter():
                    face=node.get('typeface')
                    if face in aliases:node.set('typeface',aliases[face]);changed=True
                if item.filename=='ppt/presentation.xml':
                    embedded=tree.find(P+'embeddedFontLst')
                    if embedded is not None:tree.remove(embedded);changed=True
                if changed:data=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
            out.writestr(item,data)


def audit_embedded_glyphs(source):
    """Expose actual embedded-font glyph gaps; never infer a font repair."""
    from fontTools.ttLib import TTFont
    faces={};gaps=[];checked=0;unresolved=0
    with zipfile.ZipFile(source) as z:
        pres=etree.fromstring(z.read('ppt/presentation.xml'))
        rels=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'))
        targets={x.get('Id'):x.get('Target') for x in rels}
        for embedded in pres.findall('.//'+P+'embeddedFont'):
            meta=embedded.find(P+'font')
            if meta is None:continue
            for variant in embedded:
                rid=variant.get(R+'id')
                if not rid:continue
                target=targets.get(rid,'');part=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('ppt',target))
                from .embedded_fonts import decode_font
                data,_=decode_font(z.read(part))
                font=TTFont(io.BytesIO(data));faces[(meta.get('typeface'),etree.QName(variant).localname)]=set((font.getBestCmap() or {}).keys());font.close()
        for page,part in enumerate(_parts(source),1):
            tree=etree.fromstring(z.read(part))
            for shape in tree.iter(P+'sp'):
                nv=shape.find(f'{P}nvSpPr/{P}cNvPr')
                for run in shape.findall('.//'+A+'r'):
                    text=run.findtext(A+'t',default='');props=run.find(A+'rPr')
                    if not text.strip():continue
                    if props is None:unresolved+=1;continue
                    variant='boldItalic' if props.get('b')=='1' and props.get('i')=='1' else 'bold' if props.get('b')=='1' else 'italic' if props.get('i')=='1' else 'regular'
                    slot=props.find(A+'ea') if any(ord(c)>255 for c in text) else props.find(A+'latin')
                    if slot is None:slot=props.find(A+'latin')
                    face=slot.get('typeface') if slot is not None else None
                    cmap=faces.get((face,variant),faces.get((face,'regular')))
                    if cmap is None:unresolved+=1;continue
                    checked+=1;missing=''.join(sorted({c for c in text if not c.isspace() and ord(c) not in cmap}))
                    if missing:gaps.append({'page':page,'shape_id':int(nv.get('id')) if nv is not None else None,'face':face,'text':text,'missing':missing})
    return {'embedded_faces':sorted({x[0] for x in faces}),'checked_runs':checked,'unresolved_or_system_runs':unresolved,
        'glyph_gaps':gaps,'notice':'字形缺口为诊断证据，不能单独推断整句修复；系统/继承字体及视觉风格仍需逐图检查。'}


def render_hidden(source, directory, *, timeout=600):
    from .resources import resource_slot, path_identity
    from .render_cache import environment_fingerprint, valid_cache
    source = Path(source).resolve(); directory = Path(directory).resolve()
    if not source.is_file() or source.suffix.lower() != '.pptx':
        raise ValidationError('后台渲染只接受真实PPTX')
    from .verification_environment import verification_capability, skipped_verification
    capability = verification_capability('pptx')
    if not capability['available']:
        return skipped_verification(source, capability)
    from .embedded_fonts import decoder_files
    renderer_files = [Path(__file__), Path(__file__).with_name('background_render.ps1'), *decoder_files()]
    fingerprint = environment_fingerprint(renderer_files)
    cache_file = directory/'render_cache.json'
    with resource_slot('render-output:' + path_identity(directory)):
        if cache_file.is_file():
            cached = read_json(cache_file)
            proof = cached.get('render_proof',{})
            hidden = proof.get('with_window') is False and proof.get('windows_before') == proof.get('windows_during') and proof.get('slides_exported') == len(_parts(source))
            if hidden and cached.get('foreground_opened') is False and cached.get('renderer') == RENDERER and valid_cache(source,cached,fingerprint,expected_pages=len(_parts(source))):
                return cached
        # Queue waiting is outside subprocess timeout. Existing PS mutex remains
        # the last barrier against old callers in this Windows login session.
        with resource_slot('office-background') as slot:
            result = _render_hidden(source,directory,timeout=timeout)
            if environment_fingerprint(renderer_files) != fingerprint:
                raise ValidationError('渲染期间字体或Office环境变化，旧证据不能缓存')
            result.update(environment_fingerprint=fingerprint, queue_seconds=slot['queue_seconds'])
            write_json(cache_file,result)
            return result


def _render_hidden(source, directory, *, timeout=600):
    if os.name!='nt':raise ValidationError('当前后台渲染器需要Windows及PowerPoint；禁止前台备用路线')
    shell=shutil.which('pwsh') or shutil.which('powershell')
    if not shell:raise ValidationError('缺少后台PowerShell渲染运行时')
    source=Path(source).resolve();directory=Path(directory).resolve()
    if not source.is_file() or source.suffix.lower()!='.pptx':
        raise ValidationError('后台渲染只接受现有PPTX；不启用前台转换')
    directory.mkdir(parents=True,exist_ok=True)
    digest=file_sha256(source)
    with tempfile.TemporaryDirectory(prefix='ppt_bg_fonts_') as tmp:
        fonts,aliases=extract_embedded_fonts(source,tmp)
        render_source=Path(tmp)/'render_only.pptx'
        if aliases:prepare_font_render_copy(source,render_source,aliases)
        else:render_source=source
        payload=Path(tmp)/'fonts.json';write_json(payload,fonts)
        evidence=Path(tmp)/'renderer_evidence.json'
        try:
            result=subprocess.run([shell,'-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(Path(__file__).with_name('background_render.ps1')),
                '-InputPptx',str(render_source),'-OutputDir',str(directory),'-FontsJson',str(payload),'-EvidenceJson',str(evidence)],
                capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=timeout,
                creationflags=subprocess.CREATE_NO_WINDOW)
        except subprocess.TimeoutExpired as exc:
            from .resources import mark_office_uncertain
            mark_office_uncertain(directory)
            # Only the uniquely named temporary faces owned by this render.
            # Never terminate PowerPoint or touch another user's presentation.
            import ctypes
            remove=ctypes.windll.gdi32.RemoveFontResourceExW
            for font in fonts:remove(str(font),0,None)
            raise ValidationError('后台渲染超时，已阻断；不会转为前台检查') from exc
        if result.returncode:raise ValidationError('后台渲染失败，未启用前台备用: '+(result.stderr or result.stdout)[-2500:])
        proof=json.loads(evidence.read_text(encoding='utf-8-sig'))
        if proof.get('with_window') is not False or proof.get('windows_before')!=proof.get('windows_during') or proof.get('slides_exported')!=len(_parts(source)):
            raise ValidationError('后台渲染窗口或全量导出证据不一致')
    if file_sha256(source)!=digest:raise ValidationError('后台渲染期间输入发生变化')
    from PIL import Image
    images=[]
    for n,_ in enumerate(_parts(source),1):
        path=directory/f'slide_{n:03d}.png'
        if not path.is_file():raise ValidationError('后台渲染未覆盖全部页面')
        with Image.open(path) as im:
            if im.width<100 or im.height<100:raise ValidationError('后台渲染图片尺寸异常')
            im.verify()
        images.append({'page':n,'path':str(path),'sha256':file_sha256(path)})
    return {'input_sha256':digest,'renderer':RENDERER,'foreground_opened':False,'embedded_fonts_loaded':len(fonts),
        'font_audit':audit_embedded_glyphs(source),'font_aliases_use_original_glyphs':True,'render_proof':proof,
        'images':images,'created_at':now_iso()}


def render_source_hidden(source, directory, *, force_font_fallback=False):
    from .resources import resource_slot, path_identity
    with resource_slot('source-reading:' + path_identity(Path(directory).resolve())):
        return _render_source_hidden(source, directory, force_font_fallback=force_font_fallback)


def _render_source_hidden(source, directory, *, force_font_fallback=False):
    """Content-reading entry; never used for animation or final visual acceptance."""
    from .source_font_inspection import prepare_source_copy
    source=Path(source).resolve(); directory=Path(directory).resolve()
    copy, report=prepare_source_copy(source, directory/'_font_recovery', force_fallback=force_font_fallback)
    try:
        result=dict(render_hidden(copy,directory))
    except ValidationError as exc:
        if 'Temporary embedded font registration failed' in str(exc) and not force_font_fallback:
            write_json(directory/'_font_recovery/first_render_failure.json', {'reason': str(exc)})
            return _render_source_hidden(source, directory, force_font_fallback=True)
        report.update(render_status='failed', render_failure=str(exc))
        write_json(directory/'_font_recovery/font_recovery.json',report)
        raise
    report.update(render_status='skipped' if result.get('verified') is False else 'exported',
                  embedded_fonts_loaded=result.get('embedded_fonts_loaded'),
                  render_proof=result.get('render_proof'))
    write_json(directory/'_font_recovery/font_recovery.json',report)
    result.update(input_sha256=report['source_sha256'], source_font_recovery=report,
                  original_font_appearance_verified=False, complete_content_review=False)
    if report['content_only']:
        result.update(verified=False, verification_scope='content_inspection_only')
    if file_sha256(source)!=report['source_sha256']:
        raise ValidationError('资料核对期间源文件发生变化')
    write_json(directory/'source_render_manifest.json',result)
    return result


def extract_background_material(source,directory):
    """Stage 0: ordered editable content, notes, original media, and real pages."""
    from .teaching_animation import inspect_pptx
    from pptx import Presentation
    source=Path(source).resolve();directory=Path(directory).resolve()
    info=inspect_pptx(source)
    directory=directory/info['source_sha256'].split(':')[1][:12]
    directory.mkdir(parents=True,exist_ok=True)
    prs=Presentation(source)
    for page,slide in zip(info['pages'],prs.slides):
        page['notes']=slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None else ''
    media=[];media_dir=directory/'media';media_dir.mkdir(exist_ok=True)
    with zipfile.ZipFile(source) as z:
        for name in z.namelist():
            if not name.startswith('ppt/media/') or name.endswith('/'):continue
            # Do not follow archive paths outside the designated media folder.
            target=media_dir/posixpath.basename(name)
            if any(x['name']==target.name for x in media):raise ValidationError('媒体文件名冲突')
            target.write_bytes(z.read(name));media.append({'name':target.name,'path':str(target),'sha256':file_sha256(target)})
    result={'source_sha256':info['source_sha256'],'source_path':str(source),'pages':info['pages'],'media':media,
        'complete_content_review':False,'notice':'XML和媒体已提取；必须实际逐页查看预览及图片内字，不能自动宣布完整内容审阅。','created_at':now_iso()}
    write_json(directory/'material_manifest.json',result)
    try:
        result['render']=render_source_hidden(source,directory/'previews')
    except Exception as exc:
        result['render']={'verified':False, 'status':'failed', 'reason':str(exc)}
        write_json(directory/'material_manifest.json',result)
        raise
    if file_sha256(source)!=info['source_sha256']:raise ValidationError('资料提取期间源文件发生变化')
    write_json(directory/'material_manifest.json',result)
    return result


def write_state_deck(source,target,pages):
    """Disposable render-only package; original output is never altered."""
    with zipfile.ZipFile(source) as z:
        pres=etree.fromstring(z.read('ppt/presentation.xml'));rels=etree.fromstring(z.read('ppt/_rels/presentation.xml.rels'));types=etree.fromstring(z.read('[Content_Types].xml'))
        ids=pres.find(P+'sldIdLst');ids.clear();additional={};mapping=[];index=0;token=uuid.uuid4().hex[:10]
        for page in pages:
            tree=etree.fromstring(z.read(page['part']));spmap=_top_shapes(tree)
            if tree.find(P+'timing') is not None:tree.remove(tree.find(P+'timing'))
            if tree.find(P+'transition') is not None:tree.remove(tree.find(P+'transition'))
            src_rel=posixpath.join(posixpath.dirname(page['part']),'_rels',posixpath.basename(page['part'])+'.rels')
            for state in page['states']:
                index+=1;clone=deepcopy(tree);top=_top_shapes(clone)
                for sid in state['hidden_ids']:top[sid].getparent().remove(top[sid])
                part=f'ppt/slides/bg_state_{token}_{index:04d}.xml';rid=f'rIdBgState{token}_{index}'
                if part in z.namelist() or rels.xpath('*[@Id=$id]',id=rid):raise ValidationError('临时状态包名称冲突')
                additional[part]=etree.tostring(clone,encoding='UTF-8',xml_declaration=True,standalone=True)
                if src_rel in z.namelist():additional[f'ppt/slides/_rels/bg_state_{token}_{index:04d}.xml.rels']=z.read(src_rel)
                etree.SubElement(ids,P+'sldId',{'id':str(255+index),R+'id':rid})
                etree.SubElement(rels,REL+'Relationship',Id=rid,Type=R[1:-1]+'/slide',Target=f'slides/bg_state_{token}_{index:04d}.xml')
                etree.SubElement(types,CT+'Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.presentationml.slide+xml')
                mapping.append({'render_page':index,'page':page['page'],**state})
        replaced={'ppt/presentation.xml':pres,'ppt/_rels/presentation.xml.rels':rels,'[Content_Types].xml':types}
        with zipfile.ZipFile(target,'w') as b:
            for item in z.infolist():b.writestr(item,etree.tostring(replaced[item.filename],encoding='UTF-8',xml_declaration=True) if item.filename in replaced else z.read(item.filename))
            for name,data in additional.items():b.writestr(name,data)
    return mapping


def validate_teaching_intent(entry, valid_ids):
    """Validate an independent teaching-unit contract inside the existing plan.

    The controller supplies order from visible numbering, task dependencies and
    page/media review. Never infer this expected order from the effects list.
    """
    if not isinstance(entry.get('sequence_evidence'),str) or not entry['sequence_evidence'].strip():
        raise ValidationError('每页需sequence_evidence，说明来自原页编号/推理关系的独立顺序依据；静态页说明保留理由')
    units=entry.get('teaching_units')
    if not isinstance(units,list):
        raise ValidationError('每页需teaching_units；按教学顺序声明对象、预期点击、目的及必要参照')
    seen=set();previous=0
    for unit in units:
        if not isinstance(unit,dict) or set(unit)!={'shape_ids','reveal_click','purpose','retain_ids'}:
            raise ValidationError('教学单元需且仅含shape_ids/reveal_click/purpose/retain_ids')
        ids=unit['shape_ids'];retain=unit['retain_ids'];step=unit['reveal_click']
        for name,values in [('shape_ids',ids),('retain_ids',retain)]:
            if not isinstance(values,list) or any(type(sid) is not int or sid not in valid_ids for sid in values) or len(set(values))!=len(values):
                raise ValidationError('教学单元'+name+'需互异的实际顶层对象ID')
        if not ids or seen.intersection(ids):
            raise ValidationError('教学单元对象不能为空或重复归属；组合内成员不能当作独立目标')
        if type(step) is not int or step<previous or step<0:
            raise ValidationError('教学单元须按独立预期点击顺序排列，不能由错误效果顺序反推')
        if not isinstance(unit['purpose'],str) or not unit['purpose'].strip():
            raise ValidationError('每个教学单元需具体信息或操作目的，不能用装饰凑点击')
        previous=step;seen.update(ids)
    missing = set(valid_ids) - seen
    if missing:
        raise ValidationError('teaching_units需覆盖全部实际动画对象及初始静态对象；'
                              '未归属的顶层对象ID：' + ', '.join(map(str, sorted(missing))))
    return units


def check_teaching_intent(actual, entry):
    valid=set(actual['states'][-1]['visible_ids'])
    units=validate_teaching_intent(entry,valid)
    effects={e['shape_id']:e for e in actual['effects']}
    covered={sid for u in units for sid in u['shape_ids']}
    if set(effects)-covered:
        raise ValidationError('teaching_units需覆盖全部实际动画对象，包含保留的已有动画')
    for unit in units:
        step=unit['reveal_click']
        if step>=len(actual['states']) or any(effects.get(sid,{}).get('click',0)!=step for sid in unit['shape_ids']):
            raise ValidationError(f"第{actual['page']}页实际动画与独立教学单元预期点击不一致")
        if not set(unit['retain_ids'])<=set(actual['states'][step]['visible_ids']):
            raise ValidationError('教学单元出现时缺少必要题干、例证或比较参照')
        durations=[effects[sid]['duration'] for sid in unit['shape_ids'] if sid in effects]
        if durations and max(durations)-min(durations)>.001:
            raise ValidationError('同一教学单元的同步淡入应同时完成；不要只检查同开始')


def check_visibility_rules(pages,plan):
    if len(pages)!=len(plan['slides']):raise ValidationError('后台页数与计划不一致')
    for actual,entry in zip(pages,plan['slides']):
        if actual['page']!=entry['page']:raise ValidationError('后台页序与计划不一致')
        rules=entry.get('visibility_rules')
        if entry.get('content_roles_reviewed') is not True or not isinstance(rules,list):
            raise ValidationError(f"第{entry['page']}页需主控识别题目/答案并记录visibility_rules，不能仅凭关键词推断")
        seen=set();reveal={e['shape_id']:e['click'] for e in actual['effects']};valid=set(actual['states'][-1]['visible_ids'])
        for rule in rules:
            sid=rule.get('shape_id');step=rule.get('reveal_click')
            if type(sid) is not int or sid not in valid or sid in seen or type(step) is not int or not 0<=step<len(actual['states']):raise ValidationError('可见性契约对象或点击编号无效')
            if rule.get('role') not in ('prompt','answer','explanation','illustration') or not str(rule.get('reason','')).strip():raise ValidationError('可见性规则缺少实际语义角色或理由')
            # answer means a response deliberately withheld for independent judgement.
            # Known examples/review material use prompt or explanation by teaching purpose.
            if rule['role']=='answer' and step==0:raise ValidationError('待独立判断后揭示的answer不能初始可见；已知示范或复盘材料须按实际用途标为prompt/explanation并说明理由')
            seen.add(sid)
            if reveal.get(sid,0)!=step:raise ValidationError(f"第{actual['page']}页对象{sid}的实际出现点击与语义契约不一致")
        check_teaching_intent(actual,entry)
        # Compare the serialized behavior for each new effect, including grouping
        # and duration, to the controller's plan.
        byid={e['shape_id']:e for e in actual['effects']}
        for e in entry['effects']:
            out=byid.get(e['shape_id'])
            if out is None or out['trigger']!=e['trigger'] or abs(out['duration']-e.get('duration',.45))>.001:
                raise ValidationError('成品动画与计划触发/时长不一致')
        new_ids={e['shape_id'] for e in entry['effects']}
        serialized=[e for e in actual['effects'] if e['shape_id'] in new_ids]
        if [e['shape_id'] for e in serialized]!=[e['shape_id'] for e in entry['effects']]:
            raise ValidationError('成品动画对象顺序与计划不一致')


def export_background_states(run_dir):
    from .state import read_state,write_state
    root=Path(run_dir);state=read_state(root);task=state['stage5_animation']
    if task.get('status') not in ('waiting_background_review','waiting_font_review','verification_skipped'):
        raise ValidationError('先生成并核验动画成品，再导出后台点击状态')
    for path,key in [(Path(task['input']),'input_sha256'),(Path(task['output']),'output_sha256'),(root/task['plan'],'plan_sha256')]:
        if file_sha256(path)!=task[key]:raise ValidationError('输入、输出或计划指纹变化')
    plan=read_json(root/task['plan']);pages=read_click_states(task['output']);check_visibility_rules(pages,plan)
    directory=root/'_state/阶段5/background_states'/task['output_sha256'].split(':')[1][:12];directory.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='ppt_bg_states_') as tmp:
        deck=Path(tmp)/'states.pptx';mapping=write_state_deck(task['output'],deck,pages);render=render_hidden(deck,directory)
    if render.get('method') == 'verification_skipped':
        from .verification_environment import skipped_verification
        manifest = skipped_verification(task['output'],render['capability'])
        manifest.update(input=task['input'], input_sha256=task['input_sha256'], output=task['output'], output_sha256=task['output_sha256'],
                        plan_sha256=task['plan_sha256'], pages=pages,
                        structural_timeline_checked=True, visual_states_verified=False,
                        pending_text_issues=plan.get('pending_text_issues', []),
                        text_content_accepted=not bool(plan.get('pending_text_issues')))
        task.update(status='verification_skipped', background_accepted=False, verification_skipped=True,
                    verification_skip_reason=manifest['reason'], background_manifest='_state/阶段5/background_manifest.json')
        if state['workflow']['mode'] == 'animation_only':
            if plan.get('pending_text_issues'):
                state.update(status='animation_active',required_actor='main_controller',
                             next_required_action='安全页文件与时间轴已核验；正文待办保留，缺环境未完成视觉/字体验证')
            else:
                state.update(status='completed',required_actor='none',next_required_action='动画文件与时间轴已核验；缺环境，视觉/字体验证未执行')
        from .artifact_commit import commit_project_artifact
        commit_project_artifact(root,state,task['background_manifest'],manifest)
        return manifest
    for record,image in zip(mapping,render['images']):
        record['path']=Path(image['path']).relative_to(root.resolve()).as_posix();record['sha256']=image['sha256']
    # Full final-state parity against the input preview is a meaningful extra
    # gate for timing-only work. Authorized font edits are reviewed visually.
    parity=[]
    changed_pages = {e['page'] for e in plan.get('font_repairs', []) + plan.get('text_optimizations', [])}
    if len(changed_pages) < len(pages):
        from PIL import Image, ImageChops
        previews=read_json(root/'_state/阶段5/preview_manifest.json')
        originals={x['page']:x for x in previews['images']}
        for page in pages:
            if page['page'] in changed_pages:
                parity.append({'page': page['page'], 'exact_match': None, 'method': 'authorized_text_quality_review'})
                continue
            final=next(x for x in mapping if x['page']==page['page'] and x['step']==len(page['states'])-1)
            original=root/originals[page['page']]['path']
            if file_sha256(original)!=originals[page['page']]['sha256']:
                raise ValidationError('原稿预览发生变化，不能比较最终状态')
            with Image.open(original) as a,Image.open(root/final['path']) as b:
                same=a.size==b.size and ImageChops.difference(a.convert('RGB'),b.convert('RGB')).getbbox() is None
            parity.append({'page':page['page'],'exact_match':same})
            if not same:
                raise ValidationError(f"第{page['page']}页最终后台画面与原稿不一致，需查明字体/版面/渲染差异")
    manifest={'method':'background_click_states','profile':PROFILE,'foreground_opened':False,'renderer':RENDERER,
        'environment_fingerprint':render.get('environment_fingerprint'),
        'input_sha256':task['input_sha256'],'output_sha256':task['output_sha256'],'plan_sha256':task['plan_sha256'],
        'pages':pages,'states':mapping,'embedded_fonts_loaded':render['embedded_fonts_loaded'],
        'font_audit':render.get('font_audit',{}),'render_proof':render.get('render_proof',{}),'final_frame_parity':parity,
        'scope':'完整对象短淡入的初始及每次点击完成状态；未运行目标应用放映，不保证任意播放器一致性',
        'created_at':now_iso()}
    for path,key in [(Path(task['input']),'input_sha256'),(Path(task['output']),'output_sha256'),(root/task['plan'],'plan_sha256')]:
        if file_sha256(path)!=task[key]:raise ValidationError('后台渲染期间输入、成品或计划发生变化')
    task['background_manifest']='_state/阶段5/background_manifest.json'
    if task.get('status') == 'verification_skipped':
        task.update(status='waiting_background_review',verification_skipped=False)
        if state['workflow']['mode'] == 'animation_only':
            state.update(status='animation_active',required_actor='main_controller',next_required_action='环境已恢复，完成后台状态视觉验收')
    from .artifact_commit import commit_project_artifact
    commit_project_artifact(root,state,task['background_manifest'],manifest)
    return manifest
