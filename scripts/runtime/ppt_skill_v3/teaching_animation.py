"""Independent teaching animation with explicitly authorized font compatibility repair.

Timing is written and read directly from OOXML. Visual states are rendered with
hidden document windows; no foreground slideshow is opened.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import posixpath
import shutil
import tempfile
from typing import Any
import zipfile

from lxml import etree

from .file_hashing import file_sha256
from .json_io import read_json, write_json
from .state import read_state, write_state
from .time_utils import now_iso
from .validation import ValidationError
from .font_repair import inspect_text_objects, validate_repairs, patch_font_tree, apply_font_repairs

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def slide_parts(path: str | Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        presentation = etree.fromstring(z.read("ppt/presentation.xml"))
        rels = etree.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
        mapping = {rel.get("Id"): rel.get("Target") for rel in rels if rel.get("TargetMode") != "External"}
        result = []
        for item in presentation.findall(f"{P}sldIdLst/{P}sldId"):
            target = mapping.get(item.get(f"{R}id"))
            if not target:
                raise ValidationError("slide relationship is missing")
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("ppt", target))
            if not part.startswith("ppt/slides/") or part not in z.namelist():
                raise ValidationError("invalid slide part")
            result.append(part)
        if len(set(result)) != len(result):
            raise ValidationError("duplicate slide parts")
        return result


def inspect_pptx(path: str | Path) -> dict[str, Any]:
    from pptx import Presentation
    path = Path(path).resolve()
    if not path.is_file() or path.suffix.lower() != ".pptx":
        raise ValidationError("supply an existing editable .pptx; PDF conversion is manual")
    prs = Presentation(path)
    parts = slide_parts(path)
    pages = []
    with zipfile.ZipFile(path) as z:
        for index, (slide, part) in enumerate(zip(prs.slides, parts), 1):
            root = etree.fromstring(z.read(part))
            timing = root.find(f"{P}timing")
            targets = sorted({int(n.get("spid")) for n in timing.iter(f"{P}spTgt") if n.get("spid", "").isdigit()}) if timing is not None else []
            shapes = [{"shape": i, "shape_id": s.shape_id, "name": s.name, "type": str(s.shape_type),
                       "text": s.text if s.has_text_frame else "", "has_existing_animation": s.shape_id in targets,
                       "bounds": [s.left, s.top, s.width, s.height],
                       "group_members": [c.shape_id for c in s.shapes] if hasattr(s, "shapes") else []}
                      for i, s in enumerate(slide.shapes, 1)]
            pages.append({"page": index, "part": part, "shapes": shapes,
                          "text_objects": inspect_text_objects(root),
                          "existing_targets": targets, "has_timing": timing is not None,
                          "has_transition": root.find(f"{P}transition") is not None})
    if not pages:
        raise ValidationError("empty presentation")
    return {"source_path": str(path), "source_sha256": file_sha256(path), "pages": pages,
            "created_at": now_iso(), "notice": "Inspect page previews as well as objects. Images and groups may hide question/answer pairs; do not split or rebuild them."}


def start_animation(run_dir: str | Path, input_pptx: str | Path, evidence: str, *, repair_fonts: bool = False,
                    match_reference_fonts: bool = False) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    if state['workflow']['mode'] != 'animation_only':
        from .text_preservation import verify_stage1_text
        verify_stage1_text(root,input_pptx)
    from .validation import validate_project_state_relations
    validate_project_state_relations(state)
    if state["workflow"]["mode"] != "animation_only" and state.get("status") != "completed":
        raise ValidationError("finish base delivery before starting animation; standalone animation uses its own project")
    if not evidence.strip():
        raise ValidationError("actual human request for animation is required")
    info = inspect_pptx(input_pptx)
    source = Path(info["source_path"])
    target = root / "_state/阶段5"
    target.mkdir(parents=True, exist_ok=True)
    from .font_handoff import import_font_handoffs
    handoff = import_font_handoffs(root, info, strict=match_reference_fonts)
    repair_fonts = repair_fonts or match_reference_fonts
    write_json(target / "inspection.json", info)
    state["stage5_animation"] = {"status": "planning", "input": info["source_path"],
        "input_sha256": info["source_sha256"], "authorization_evidence": evidence,
        "font_repair_enabled": repair_fonts, "font_review_accepted": not repair_fonts,
        "reference_font_matching_enabled": match_reference_fonts,
        "reference_font_authorization_evidence": evidence if match_reference_fonts else None,
        "font_handoff_import": "_state/阶段5/font_handoff_import.json",
        "font_handoff_import_sha256": file_sha256(target / 'font_handoff_import.json'),
        "font_handoff_pending_count": len(handoff['items']),
        "inspection": "_state/阶段5/inspection.json", "background_accepted": False,
        "execution_mode": "background_only"}
    if state['workflow']['mode'] != 'animation_only' and any(
            p['match_method'] == 'controller_reviewed_display_mapping' for p in verify_stage1_text(root,input_pptx)['pages']):
        from .text_mapping_review import REVIEW
        state['stage5_animation'].update({'text_mapping_review': REVIEW.as_posix(),
            'text_mapping_review_sha256': file_sha256(root / REVIEW)})
    if state["workflow"]["mode"] == "animation_only":
        state.update({"status": "animation_active", "required_actor": "main_controller", "next_required_action": "规划当前输入版本的教学动画"})
    write_state(root, state)
    return info


def record_animation_plan(run_dir: str | Path, plan: dict[str, Any]) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    task = state["stage5_animation"]
    from .text_mapping_review import require_task_mapping
    require_task_mapping(root, task)
    previous_status = task.get('status')
    if previous_status not in {"planning", "needs_manual_adjustment", "failed", "plan_ready", "verification_skipped"}:
        raise ValidationError("start inspection before recording the animation plan")
    info = read_json(root / "_state/阶段5/inspection.json")
    if file_sha256(task["input"]) != info["source_sha256"] or plan.get("input_sha256") != info["source_sha256"]:
        raise ValidationError("PPTX changed; re-inspect and re-map the new input")
    if plan.get("controller_reviewed") is not True or not plan.get("preview_review_evidence"):
        raise ValidationError("controller must inspect page previews and supply concrete review evidence")
    _require_previews(root, task)
    preview = read_json(root / '_state/阶段5/preview_manifest.json')
    visual_skipped = preview.get('method') == 'verification_skipped'
    slides = plan.get("slides")
    if not isinstance(slides, list) or len(slides) != len(info["pages"]):
        raise ValidationError("one plan entry per actual presentation page is required")
    seen = set()
    for entry in slides:
        page = entry.get("page")
        if type(page) is not int or page in seen or not 1 <= page <= len(info["pages"]):
            raise ValidationError("invalid/duplicate animation page")
        seen.add(page)
        actual = info["pages"][page - 1]
        for key in ("teaching_reason", "initial_state", "final_state"):
            if not isinstance(entry.get(key), str) or not entry[key].strip():
                raise ValidationError(f"animation page needs {key}; zero effects also needs a reason")
        if not isinstance(entry.get("manual_issues"), list):
            raise ValidationError("manual_issues must be a list")
        if entry.get("content_roles_reviewed") is not True or not isinstance(entry.get("visibility_rules"), list):
            raise ValidationError("each page needs reviewed content roles and visibility_rules for background answer checks")
        effects, replace = entry.get("effects"), entry.get("replace_targets", [])
        if not isinstance(effects, list) or not isinstance(replace, list) or len(replace) != len(set(replace)):
            raise ValidationError("effects and distinct replace_targets are required")
        valid_shapes = {s["shape"]: s for s in actual["shapes"]}
        for shape in replace:
            if type(shape) is not int or shape not in valid_shapes:
                raise ValidationError("replacement target does not exist")
            if not valid_shapes[shape]["has_existing_animation"]:
                raise ValidationError("replacement target has no existing animation")
        used = set()
        for effect in effects:
            shape = effect.get("shape")
            if type(shape) is not int or shape not in valid_shapes or shape in used:
                raise ValidationError("animation target must be a distinct existing top-level object")
            used.add(shape)
            if effect.get("shape_id") != valid_shapes[shape]["shape_id"]:
                raise ValidationError("shape identity changed")
            if valid_shapes[shape]["has_existing_animation"] and shape not in replace:
                raise ValidationError("target already animated: preserve or explicitly replace its effects")
            if effect.get("trigger") not in {"click", "with_previous"}:
                raise ValidationError("unsupported trigger")
            duration = effect.get("duration", 0.45)
            if isinstance(duration, bool) or not isinstance(duration, (int, float)) or not 0.2 <= duration <= 1.2:
                raise ValidationError("use a short fade duration (0.2–1.2 seconds)")
            for key in ("role", "reason", "click_group"):
                if not isinstance(effect.get(key), str) or not effect[key].strip():
                    raise ValidationError(f"effect needs {key}")
            if effect.get("separable") is not True:
                raise ValidationError("mixed prompt/answer objects must stay unchanged and become manual issues")
        group = None
        group_names = set()
        for effect in effects:
            if effect["trigger"] == "click":
                if effect["click_group"] in group_names:
                    raise ValidationError("distinct clicks must have distinct click_group names")
                group = effect["click_group"]
                group_names.add(group)
            elif group is None or effect["click_group"] != group:
                raise ValidationError("with_previous must follow its own click group")
        if set(replace) - used:
            raise ValidationError("replacement targets need explicit replacement effects; do not erase arbitrary timing")
    repairs = plan.get("font_repairs", [])
    if visual_skipped and repairs:
        raise ValidationError('without current PPTX visual verification preserve fonts and record unverified issues')
    validate_repairs(Path(task["input"]), slide_parts(task["input"]), repairs)
    if repairs and not task.get("font_repair_enabled"):
        raise ValidationError("font repair was not authorized; pure animation preserves every font")
    _require_reference_authorization(task, repairs)
    from .font_handoff import validate_handoff_resolutions
    validate_handoff_resolutions(root, task, plan, info)
    if task.get("font_repair_enabled") and (not visual_skipped or repairs):
        if not isinstance(plan.get("font_review_evidence"), str) or not plan["font_review_evidence"].strip():
            raise ValidationError("actually inspect all page fonts, even when no repairs are needed")
        for entry in slides:
            if not isinstance(entry.get("font_observation"), str) or not entry["font_observation"].strip():
                raise ValidationError("each page needs an actual font observation")
        _bind_font_evidence(root, repairs, bind=True)
        _bind_font_evidence(root, plan.get('font_handoff_resolutions', []), bind=True)
    result = {**plan, "font_repairs": repairs, "created_at": now_iso()}
    if visual_skipped:
        result['visual_evidence_status'] = 'skipped_no_environment'
        result['visual_verification_reason'] = preview['reason']
    else:
        result.pop('visual_evidence_status', None)
        result.pop('visual_verification_reason', None)
    write_json(root / "_state/阶段5/animation_plan.json", result)
    task["status"] = "needs_manual_adjustment" if any(s["manual_issues"] for s in slides) else "plan_ready"
    task["plan"] = "_state/阶段5/animation_plan.json"
    task["plan_sha256"] = file_sha256(root / task["plan"])
    task["preview_manifest_sha256"] = file_sha256(root / "_state/阶段5/preview_manifest.json")
    if previous_status == 'verification_skipped' and state['workflow']['mode'] == 'animation_only':
        state.update(status='animation_active', required_actor='main_controller')
    write_state(root, state)
    return result


def _bind_font_evidence(root: Path, repairs: list[dict], *, bind: bool) -> None:
    for repair in repairs:
        if repair.get('status') == 'blocked':
            continue
        from .font_handoff import local_evidence
        keys = ['visual_evidence']
        if 'mode' in repair or repair.get('font_file'):
            keys += ['sample_review_evidence', 'font_file']
        if repair.get('mode') == 'reference_match':
            keys += ['reference_visual']
        for key in keys:
            path = local_evidence(root, repair.get(key))
            digest = file_sha256(path)
            if bind:
                repair[key + '_sha256'] = digest
            elif digest != repair.get(key + '_sha256'):
                raise ValidationError('font diagnosis/reference/font file changed; review the plan again')


def _require_reference_authorization(task, repairs):
    if any(r.get('mode') == 'reference_match' for r in repairs) and not task.get('reference_font_matching_enabled'):
        raise ValidationError('reference font matching requires explicit separate authorization')


def export_animation_previews(run_dir: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    task = state["stage5_animation"]
    if not task.get("input") or file_sha256(task["input"]) != task.get("input_sha256"):
        raise ValidationError("inspect the current input before rendering previews")
    from .background_ppt import render_hidden
    directory = root.resolve() / "_state/阶段5/previews" / task["input_sha256"].split(":")[1][:12]
    manifest = render_hidden(task["input"], directory)
    for image in manifest["images"]:
        image["path"] = Path(image["path"]).relative_to(root.resolve()).as_posix()
    task["previews"] = "_state/阶段5/preview_manifest.json"
    from .artifact_commit import commit_project_artifact
    commit_project_artifact(root, state, task['previews'], manifest)
    return manifest


def _require_previews(root: Path, task: dict[str, Any]) -> None:
    path = root / "_state/阶段5/preview_manifest.json"
    if not path.is_file():
        raise ValidationError("export and actually inspect every current PPTX page preview before planning")
    manifest = read_json(path)
    if manifest.get('method') == 'verification_skipped':
        from .verification_environment import verification_capability
        if verification_capability('pptx')['available'] or manifest.get('input_sha256') != task['input_sha256'] or manifest.get('verified') is not False:
            raise ValidationError('后台环境已具备或来源已变，必须重新导出预览')
        return
    if manifest.get('environment_fingerprint'):
        from .render_cache import environment_fingerprint
        renderer_module = Path(__file__).with_name('background_ppt.py')
        if manifest['environment_fingerprint'] != environment_fingerprint([renderer_module,renderer_module.with_name('background_render.ps1')]):
            raise ValidationError('字体或Office渲染环境已变化，须重新预览和制定计划')
    info = read_json(root / "_state/阶段5/inspection.json")
    images = manifest.get("images", [])
    if manifest.get("input_sha256") != task["input_sha256"] or len(images) != len(info["pages"]) or {i.get("page") for i in images} != {p["page"] for p in info["pages"]}:
        raise ValidationError("previews do not cover the current input")
    for image in images:
        actual = root / image["path"]
        if not actual.is_file() or file_sha256(actual) != image.get("sha256"):
            raise ValidationError("preview image changed or is missing")


def merge_timing(source: Path, animated: Path, output: Path, pages: set[int]) -> None:
    source_parts, animated_parts = slide_parts(source), slide_parts(animated)
    if len(source_parts) != len(animated_parts):
        raise ValidationError("PowerPoint changed slide count")
    part_map = {source_parts[p - 1]: animated_parts[p - 1] for p in pages}
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(animated) as changed, zipfile.ZipFile(output, "w") as out:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename in part_map:
                before = etree.fromstring(data)
                after = etree.fromstring(changed.read(part_map[info.filename]))
                ids = lambda r: [n.get("id") for n in r.iter(f"{P}cNvPr")]
                if ids(before) != ids(after):
                    raise ValidationError("PowerPoint changed shape identities")
                timing = after.find(f"{P}timing")
                if timing is None:
                    raise ValidationError("PowerPoint did not save requested timing")
                existing = before.find(f"{P}timing")
                if existing is not None:
                    before.remove(existing)
                ext = before.find(f"{P}extLst")
                before.insert(before.index(ext) if ext is not None else len(before), deepcopy(timing))
                data = etree.tostring(before, encoding="UTF-8", xml_declaration=True, standalone=True)
            out.writestr(info, data)


def verify_animation_only(source: Path, output: Path, pages: set[int], font_repairs: list[dict] | None = None) -> dict[str, Any]:
    from pptx import Presentation
    repairs = font_repairs if font_repairs is not None else []
    parts = slide_parts(source)
    validate_repairs(source, parts, repairs)
    timing_allowed = {parts[p - 1] for p in pages}
    fonts = {}
    for repair in repairs:
        fonts.setdefault(parts[repair["page"] - 1], []).append(repair)
    allowed = timing_allowed | set(fonts)
    changed = []
    changed_fonts = []
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(output) as b:
        if a.namelist() != b.namelist():
            raise ValidationError("package entries/order changed")
        for name in a.namelist():
            before, after = a.read(name), b.read(name)
            if before == after and name not in fonts:
                continue
            if name not in allowed:
                raise ValidationError(f"non-animation package change: {name}")
            roots = [etree.fromstring(before), etree.fromstring(after)]
            if name in fonts:
                patch_font_tree(roots[0], fonts[name])
                changed_fonts.append(name)
            if name in timing_allowed:
                original_timing = roots[0].find(f"{P}timing")
                output_timing = roots[1].find(f"{P}timing")
                canonical = lambda t: etree.tostring(t, method="c14n") if t is not None else None
                if canonical(original_timing) != canonical(output_timing):
                    changed.append(name)
                for tree in roots:
                    timing = tree.find(f"{P}timing")
                    if timing is not None:
                        tree.remove(timing)
            if etree.tostring(roots[0], method="c14n") != etree.tostring(roots[1], method="c14n"):
                raise ValidationError(f"XML changed outside planned timing/font properties: {name}")
            shape_ids = {int(n.get("id")) for n in etree.fromstring(after).iter(f"{P}cNvPr")}
            for target in etree.fromstring(after).iter(f"{P}spTgt"):
                if not target.get("spid", "").isdigit() or int(target.get("spid")) not in shape_ids:
                    raise ValidationError("invalid timing object reference")
    if slide_parts(source) != slide_parts(output) or len(Presentation(output).slides) != len(slide_parts(source)):
        raise ValidationError("presentation order changed or final file is unreadable")
    return {"non_timing_changes": repairs, "allowed_font_changes": repairs,
            "changed_font_parts": changed_fonts, "unexpected_changes": [],
            "changed_timing_parts": changed, "openable": True,
            "source_sha256": file_sha256(source), "output_sha256": file_sha256(output)}


def _unique_output(directory: Path, stem: str) -> Path:
    candidate = directory / f"{stem}_可编辑动画版.pptx"
    version = 2
    while candidate.exists():
        candidate = directory / f"{stem}_可编辑动画版_v{version}.pptx"
        version += 1
    return candidate


def execute_animation(run_dir: str | Path) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    task = state["stage5_animation"]
    from .text_mapping_review import require_task_mapping
    require_task_mapping(root, task)
    if task.get("status") != "plan_ready":
        raise ValidationError("a reviewed plan with no unresolved manual issues is required")
    source = Path(task["input"])
    plan_path = root / task["plan"]
    if file_sha256(source) != task["input_sha256"] or file_sha256(plan_path) != task["plan_sha256"]:
        raise ValidationError("source or plan changed: inspect/review again")
    _require_previews(root, task)
    if file_sha256(root / "_state/阶段5/preview_manifest.json") != task.get("preview_manifest_sha256"):
        raise ValidationError("preview basis changed after the plan was reviewed")
    plan = read_json(plan_path)
    repairs = plan.get("font_repairs", [])
    if repairs and not task.get("font_repair_enabled"):
        raise ValidationError("font repair was not authorized")
    _require_reference_authorization(task, repairs)
    validate_repairs(source, slide_parts(source), repairs)
    from .font_handoff import validate_handoff_resolutions
    validate_handoff_resolutions(root, task, plan, read_json(root / '_state/阶段5/inspection.json'))
    _bind_font_evidence(root, repairs, bind=False)
    _bind_font_evidence(root, plan.get('font_handoff_resolutions', []), bind=False)
    pages = {s["page"] for s in plan["slides"] if s["effects"] or s.get("replace_targets")}
    task["status"] = "executing"
    task.pop("error", None)
    write_state(root, state)
    try:
        with tempfile.TemporaryDirectory(prefix="ppt_skill_animation_") as tmp:
            temp = Path(tmp)
            merged = temp / "merged.pptx"
            from .background_ppt import write_timing_file, read_click_states, check_visibility_rules
            write_timing_file(source, merged, plan)
            if repairs:
                repaired = temp / "font_repaired.pptx"
                apply_font_repairs(source, merged, repaired, slide_parts(source), repairs)
                merged = repaired
            report = verify_animation_only(source, merged, pages, repairs)
            check_visibility_rules(read_click_states(merged), plan)
            if file_sha256(source) != task["input_sha256"]:
                raise ValidationError("source changed during animation execution")
            from .resources import resource_slot, path_identity
            with resource_slot('output-directory:' + path_identity(source.parent)):
                output = _unique_output(source.parent, source.stem)
                # Exclusive reservation also protects against an unrelated writer.
                with output.open('xb') as destination, merged.open('rb') as generated:
                    shutil.copyfileobj(generated, destination)
        report.update({"output": str(output), "effects": sum(len(s["effects"]) for s in plan["slides"]),
                       "font_repair_enabled": task.get("font_repair_enabled", False),
                       "font_review_accepted": not task.get("font_repair_enabled", False),
                       "plan_sha256": task["plan_sha256"], "background_accepted": False,
                       "execution_mode": "background_only", "created_at": now_iso()})
        write_json(root / "_state/阶段5/file_verification.json", report)
        task.update({"status": "waiting_background_review", "output": str(output),
                     "background_accepted": False,
                     "font_review_accepted": not task.get("font_repair_enabled", False),
                     "output_sha256": report["output_sha256"], "file_verification": "_state/阶段5/file_verification.json"})
        write_state(root, state)
        return report
    except Exception as exc:
        task.update({"status": "failed", "error": str(exc)})
        write_state(root, state)
        raise


def accept_slideshow_review(run_dir: str | Path, review: dict[str, Any]) -> dict[str, Any]:
    raise ValidationError("foreground slideshow acceptance was removed; use export-animation-states and record-animation-background-review")


def validate_background_review(root: Path, task: dict[str, Any], review: dict[str, Any]) -> dict[str, Any]:
    from .background_ppt import read_click_states, check_visibility_rules, PROFILE, RENDERER
    if review.get("method") != "background_click_states" or review.get("controller_reviewed") is not True:
        raise ValidationError("background_click_states controller review is required")
    for key, path in (("input_sha256", Path(task["input"])), ("output_sha256", Path(task["output"])), ("plan_sha256", root/task["plan"])):
        if review.get(key) != task[key] or file_sha256(path) != task[key]:
            raise ValidationError("background review is not bound to current input/output/plan")
    manifest_path = root/task.get("background_manifest", "_state/阶段5/background_manifest.json")
    digest = file_sha256(manifest_path)
    if digest != task.get("background_manifest_sha256") or review.get("manifest_sha256") != digest:
        raise ValidationError("background render manifest changed or is unbound")
    manifest = read_json(manifest_path)
    if manifest.get('environment_fingerprint'):
        from .render_cache import environment_fingerprint
        renderer_module = Path(__file__).with_name('background_ppt.py')
        if manifest['environment_fingerprint'] != environment_fingerprint([renderer_module,renderer_module.with_name('background_render.ps1')]):
            raise ValidationError('后台状态渲染环境已变化，须重新导出和审阅')
    if manifest.get("profile") != PROFILE or manifest.get("foreground_opened") is not False or manifest.get("renderer") != RENDERER:
        raise ValidationError("unsupported or foreground renderer evidence")
    _require_previews(root, task)
    if file_sha256(root/"_state/阶段5/preview_manifest.json") != task.get("preview_manifest_sha256"):
        raise ValidationError("original preview basis changed after planning")
    for key in ("input_sha256", "output_sha256", "plan_sha256"):
        if manifest.get(key) != task[key]: raise ValidationError("stale background manifest")
    actual = read_click_states(task["output"])
    plan = read_json(root/task["plan"])
    check_visibility_rules(actual, plan)
    verify_animation_only(Path(task["input"]),Path(task["output"]),
        {s["page"] for s in plan["slides"] if s["effects"] or s.get("replace_targets")},plan.get("font_repairs",[]))
    if actual != manifest.get("pages"):
        raise ValidationError("background states do not match the final PPTX timing")
    expected = {(p["page"],s["step"]):s for p in actual for s in p["states"]}
    proof = manifest.get("render_proof", {})
    if proof.get("with_window") is not False or type(proof.get("windows_before")) is not int or type(proof.get("windows_during")) is not int or proof.get("windows_before",-1)<0 or proof.get("windows_before") != proof.get("windows_during") or proof.get("slides_exported") != len(expected):
        raise ValidationError("background review requires hidden-window and full export evidence")
    records = manifest.get("states", [])
    if len(records) != len(expected) or {(r.get("page"),r.get("step")) for r in records} != set(expected):
        raise ValidationError("render manifest omits or duplicates click states")
    images = {}
    for rec in records:
        key=(rec["page"],rec["step"])
        if any(rec.get(k)!=expected[key][k] for k in ("visible_ids","hidden_ids")):
            raise ValidationError("render state visibility was changed")
        image=root/rec["path"]
        if not image.is_file() or file_sha256(image)!=rec["sha256"]:
            raise ValidationError("background state image is missing or changed")
        images[key]=rec
    observations=review.get("states", [])
    if len(observations)!=len(expected) or {(r.get("page"),r.get("step")) for r in observations}!=set(expected):
        raise ValidationError("review must cover every initial and settled click state exactly once")
    for obs in observations:
        rec=images[(obs["page"],obs["step"])]
        if obs.get("pass") is not True or obs.get("image_sha256")!=rec["sha256"]:
            raise ValidationError("click state has unresolved issues or stale evidence")
        for key in ("visibility_observation","answer_leakage_observation","background_residue_observation","font_observation","layout_observation"):
            if not isinstance(obs.get(key),str) or not obs[key].strip():
                raise ValidationError("each state needs concrete "+key)
    if review.get("unresolved_issues") != []:
        raise ValidationError("background review requires an explicit empty unresolved_issues list")
    return manifest


def accept_background_review(run_dir: str | Path, review: dict[str, Any]) -> dict[str, Any]:
    root = Path(run_dir)
    state = read_state(root)
    task = state["stage5_animation"]
    if task.get("status") not in {"waiting_background_review", "waiting_font_review"}:
        raise ValidationError("generate and verify output before background acceptance")
    validate_background_review(root,task,review)
    plan = read_json(root / task['plan'])
    _bind_font_evidence(root, plan.get('font_repairs', []), bind=False)
    _bind_font_evidence(root, plan.get('font_handoff_resolutions', []), bind=False)
    from .font_handoff import validate_handoff_resolutions
    validate_handoff_resolutions(root, task, plan, read_json(root / '_state/阶段5/inspection.json'))
    write_json(root/"_state/阶段5/background_review.json",{**review,"created_at":now_iso()})
    task["background_accepted"] = True
    task["background_review_sha256"] = file_sha256(root/"_state/阶段5/background_review.json")
    # Font/layout observations cover every state, including all final pages.
    # Store a distinct font evidence file without re-opening any application.
    if task.get("font_repair_enabled"):
        write_json(root/"_state/阶段5/font_review.json",{**review,"method":"background_page_render","created_at":now_iso()})
        task["font_review_accepted"]=True
        task["font_review_application"]="PowerPoint hidden Slide.Export (not actual slideshow or WPS)"
    _complete_reviews(state)
    if task.get('status') == 'completed' and task.get('reference_font_matching_enabled'):
        plan = read_json(root / task['plan'])
        write_json(root / '_state/阶段5/font_handoff_results.json', {
            'input_sha256': task['input_sha256'], 'output_sha256': task['output_sha256'],
            'plan_sha256': task['plan_sha256'], 'background_review_sha256': task['background_review_sha256'],
            'scope': 'exported_pptx_only_canva_design_unchanged',
            'items': plan.get('font_handoff_resolutions', []), 'verified': True, 'created_at': now_iso()})
    write_state(root, state)
    return task


def _complete_reviews(state: dict[str, Any]) -> None:
    task = state["stage5_animation"]
    fonts_done = not task.get("font_repair_enabled") or task.get("font_review_accepted") is True
    complete = fonts_done and task.get("background_accepted") is True
    task["status"] = "completed" if complete else ("waiting_font_review" if task.get("background_accepted") else "waiting_background_review")
    if state["workflow"]["mode"] == "animation_only":
        state.update({"status": "completed" if complete else "animation_active",
                      "required_actor": "none" if complete else "main_controller",
                      "next_required_action": "动画专项后台验收已完成" if complete else "完成后台点击状态、字体与排版复核"})


def accept_font_review(run_dir: str | Path, review: dict[str, Any]) -> dict[str, Any]:
    """A font entry cannot bypass the complete background state evidence gate."""
    if review.get("method") != "background_page_render":
        raise ValidationError("font acceptance requires the complete background manifest and state observations")
    return accept_background_review(run_dir, {**review, "method":"background_click_states"})
